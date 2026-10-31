from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import pandas as pd
import pandera.pandas as pa
from openpyxl import load_workbook


class InvalidFile(ValueError):
    pass


def identifier(value: str, fallback: str = "dataset") -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    value = re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_") or fallback
    if value[0].isdigit():
        value = "c_" + value
    return value[:63].rstrip("_")


def normalize_headers(headers: list[Any]) -> list[str]:
    names, used = [], set()
    for position, header in enumerate(headers, 1):
        base = identifier(str(header) if header is not None else "", f"column_{position}")
        name, suffix = base, 1
        while name in used:
            suffix += 1
            tail = f"_{suffix}"
            name = base[:63 - len(tail)] + tail
        names.append(name)
        used.add(name)
    return names


@dataclass
class Parsed:
    headers: list[str]
    names: list[str]
    rows: list[tuple[int, list[Any]]]
    received: int
    errors: list[dict] = field(default_factory=list)


def error(row: int | None, column: str | None, stage: str, message: str) -> dict:
    return {"row_number": row, "column_name": column, "stage": stage, "message": message}


def parse_file(content: bytes, filename: str, max_rows: int, max_columns: int) -> Parsed:
    extension = Path(filename).suffix.lower()
    workbook = None
    try:
        if extension == ".csv":
            # The csv module retains lexical values, including leading zero IDs.
            stream = io.StringIO(content.decode("utf-8-sig"), newline="")
            reader = iter(csv.reader(stream, strict=True))
        elif extension == ".xlsx":
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                    raise InvalidFile("Expanded Excel workbook exceeds 200 MiB.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet = workbook.worksheets[0]
            if sheet.max_column and sheet.max_column > max_columns:
                raise InvalidFile("Too many columns in the worksheet.")

            def excel_rows():
                for cells in sheet.iter_rows():
                    values = []
                    for cell in cells:
                        if cell.data_type in {"f", "e"}:
                            values.append(InvalidCell("Excel formulas/errors are not accepted."))
                        elif (
                            isinstance(cell.value, (int, float))
                            and not isinstance(cell.value, bool)
                            and cell.value % 1 == 0
                            and re.fullmatch(r"0+", cell.number_format)
                        ):
                            # Preserve Excel's displayed zero-padded identifier format.
                            values.append(f"{int(cell.value):0{len(cell.number_format)}d}")
                        else:
                            values.append(cell.value)
                    yield values

            reader = iter(excel_rows())
        else:
            raise InvalidFile("Upload a .csv or .xlsx file.")
        header = next(reader, None)
        if not header or not any(value is not None and str(value).strip() for value in header):
            raise InvalidFile("A non-empty header row is required.")
        if len(header) > max_columns:
            raise InvalidFile("Too many columns.")
        if any(isinstance(value, InvalidCell) for value in header):
            raise InvalidFile("Headers cannot contain Excel formulas/errors.")
        parsed = Parsed([str(x) if x is not None else "" for x in header], normalize_headers(header), [], 0)
        for row_number, row in enumerate(reader, 2):
            parsed.received += 1
            if parsed.received > max_rows:
                raise InvalidFile(f"File exceeds {max_rows} data rows.")
            if len(row) != len(header):
                parsed.errors.append(error(row_number, None, "PARSE", "Row has a different number of columns than the header."))
            else:
                parsed.rows.append((row_number, list(row)))
        if not parsed.received:
            raise InvalidFile("The file has no data rows.")
        return parsed
    except InvalidFile:
        raise
    except (UnicodeError, csv.Error, zipfile.BadZipFile, OSError, ValueError, KeyError, IndexError, SyntaxError) as exc:
        raise InvalidFile("Cannot parse file. Use UTF-8 CSV or a valid .xlsx workbook.") from exc
    finally:
        if workbook is not None:
            workbook.close()


@dataclass
class InvalidCell:
    message: str


def clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def convert(value: Any, kind: str) -> Any:
    if value is None:
        return None
    if isinstance(value, InvalidCell):
        raise ValueError(value.message)
    text = str(value)
    if "\x00" in text:
        raise ValueError("NUL characters are not supported by PostgreSQL.")
    if kind == "TEXT":
        return text
    if kind == "BOOLEAN":
        if text.lower() not in {"true", "false", "yes", "no"}:
            raise ValueError("Expected true/false or yes/no.")
        return text.lower() in {"true", "yes"}
    if kind in {"INTEGER", "BIGINT"}:
        if not re.fullmatch(r"[+-]?\d+", text):
            raise ValueError("Expected an integer.")
        number = int(text)
        bound = 2**31 if kind == "INTEGER" else 2**63
        if not -bound <= number < bound:
            raise ValueError(f"Value is outside {kind} range.")
        return number
    if kind == "NUMERIC":
        number = Decimal(text)
        if not number.is_finite() or abs(number.adjusted()) > 1000 or len(number.as_tuple().digits) > 1000:
            raise ValueError("Expected a finite decimal with at most 1000 digits/exponent.")
        return number
    if kind == "DATE":
        if isinstance(value, datetime):
            if value.time() != datetime.min.time():
                raise ValueError("Expected a date without time.")
            return value.date()
        return date.fromisoformat(text)
    if kind == "TIMESTAMP":
        if isinstance(value, datetime):
            result = value
        else:
            if "T" not in text and " " not in text:
                raise ValueError("Expected ISO timestamp.")
            result = datetime.fromisoformat(text.replace("Z", "+00:00"))
        if result.tzinfo is not None:
            raise ValueError("Timezone-aware timestamps require an explicit timezone policy.")
        return result
    raise ValueError("Unknown inferred type.")


def infer_type(name: str, values: list[Any]) -> str:
    values = [v for v in values if v is not None and not isinstance(v, InvalidCell)]
    if not values:
        return "TEXT"
    # IDs and padded numeric strings are never interpreted as quantities.
    if name == "id" or name.endswith("_id") or name in {"zip", "zip_code", "postal_code", "phone"}:
        return "TEXT"
    if any(isinstance(v, str) and re.fullmatch(r"[+-]?0\d+", v) for v in values):
        return "TEXT"
    for kind in ("BOOLEAN", "INTEGER", "BIGINT", "NUMERIC", "DATE", "TIMESTAMP"):
        successes = 0
        for value in values:
            try:
                convert(value, kind)
                successes += 1
            except (ValueError, TypeError, InvalidOperation, OverflowError):
                pass
        # Allow a small minority of corrupt cells to become row errors.
        if successes / len(values) >= 0.95:
            if kind in {"INTEGER", "BIGINT"}:
                fractional = False
                for value in values:
                    try:
                        number = convert(value, "NUMERIC")
                        if number != number.to_integral_value():
                            fractional = True
                    except (ValueError, TypeError, InvalidOperation, OverflowError):
                        pass
                if fractional:
                    continue
            # Integer overflow must widen, never reject valid larger numbers.
            if kind in {"INTEGER", "BIGINT"}:
                bound = 2**(31 if kind == "INTEGER" else 63)
                out_of_range = False
                for value in values:
                    if re.fullmatch(r"[+-]?\d+", str(value)):
                        try:
                            out_of_range |= not -bound <= int(str(value)) < bound
                        except ValueError:
                            # Python limits exceptionally long integer strings;
                            # conversion records these as invalid cells later.
                            pass
                if out_of_range:
                    continue
            return kind
    return "TEXT"


def profile(values: list[Any], kind: str) -> dict:
    nonnull = [v for v in values if v is not None]
    strings = [str(v) for v in nonnull]
    result = {
        "null_count": len(values) - len(nonnull), "unique_count": len(set(strings)),
        "sample_values": list(dict.fromkeys(strings))[:5],
    }
    numbers = []
    if kind in {"INTEGER", "BIGINT", "NUMERIC"}:
        for value in nonnull:
            try:
                numbers.append(Decimal(str(convert(value, kind))))
            except (ValueError, TypeError, InvalidOperation, OverflowError):
                pass
    if numbers:
        result.update(min=str(min(numbers)), max=str(max(numbers)), mean=str(sum(numbers) / len(numbers)))
    return result


def validate_raw(parsed: Parsed, schema: list[dict]) -> tuple[list, list[dict]]:
    """Validate lexical input against inferred types before materializing conversions."""
    if not parsed.rows:
        return [], []

    def valid(value, kind):
        try:
            convert(clean(value), kind)
            return True
        except (ValueError, TypeError, InvalidOperation, OverflowError):
            return False

    validation = pa.DataFrameSchema({
        col["name"]: pa.Column(object, nullable=True, checks=pa.Check(
            lambda value, kind=col["data_type"]: valid(value, kind), element_wise=True,
        )) for col in schema
    }, strict=True)
    frame = pd.DataFrame([row for _, row in parsed.rows], columns=parsed.names, dtype=object)
    errors, bad = [], set()
    kinds = {col["name"]: col["data_type"] for col in schema}
    try:
        validation.validate(frame, lazy=True)
    except pa.errors.SchemaErrors as exc:
        for failure in exc.failure_cases.to_dict("records"):
            index = failure.get("index")
            if index is None or pd.isna(index):
                raise ValueError("Input schema failed validation.") from exc
            index = int(index)
            bad.add(index)
            name = failure["column"]
            errors.append(error(parsed.rows[index][0], name, "VALIDATE", f"Invalid {kinds[name]} value."))
    return [row for index, row in enumerate(parsed.rows) if index not in bad], errors


def prepare(parsed: Parsed) -> tuple[list[dict], list[tuple[int, dict]], list[dict]]:
    # Trimmed projections are used only for inference/profiling; raw rows are
    # retained until the first validation pass finishes.
    projected = [(number, [clean(value) for value in row]) for number, row in parsed.rows]
    schema = []
    for index, name in enumerate(parsed.names):
        values = [row[index] for _, row in projected]
        kind = infer_type(name, values)
        schema.append({"name": name, "original_name": parsed.headers[index], "data_type": kind,
                       "nullable": True, "profile": profile(values, kind)})
    validated, validation_errors = validate_raw(parsed, schema)
    errors, records = [*parsed.errors, *validation_errors], []
    for number, raw_row in validated:
        row = [clean(value) for value in raw_row]
        record, invalid = {}, False
        for column, value in zip(schema, row):
            try:
                record[column["name"]] = convert(value, column["data_type"])
            except (ValueError, TypeError, InvalidOperation, OverflowError):
                errors.append(error(number, column["name"], "TRANSFORM", f"Invalid {column['data_type']} value."))
                invalid = True
        if not invalid:
            records.append((number, record))

    # Re-validate the transformed Python objects without pandas' float/NaN coercion.
    expected = {"TEXT": str, "INTEGER": int, "BIGINT": int, "NUMERIC": Decimal,
                "BOOLEAN": bool, "DATE": date, "TIMESTAMP": datetime}
    validation = pa.DataFrameSchema({
        col["name"]: pa.Column(
            object, nullable=True,
            checks=pa.Check(lambda series, typ=expected[col["data_type"]]: series.map(lambda v: isinstance(v, typ))),
        ) for col in schema
    }, strict=True)
    if records:
        frame = pd.DataFrame([row for _, row in records], dtype=object)
        try:
            validation.validate(frame, lazy=True)
        except pa.errors.SchemaErrors as exc:
            bad = set()
            for failure in exc.failure_cases.to_dict("records"):
                index = failure.get("index")
                if index is None or pd.isna(index):
                    raise ValueError("Transformed schema failed validation.") from exc
                index = int(index)
                bad.add(index)
                errors.append(error(records[index][0], str(failure["column"]), "VALIDATE", "Value failed inferred schema validation."))
            records = [record for index, record in enumerate(records) if index not in bad]
    return schema, records, errors
